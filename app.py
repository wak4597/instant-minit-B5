import streamlit as st
import docx
from docxtpl import DocxTemplate, RichText
import re
import io
import zipfile

# Custom list of English words/phrases to italicize
ENGLISH_WORDS = [
    "Mobile Dialysis Machine", 
    "Negative Pressure Isolation Room", 
    "Treatment Room", 
    "Consultation Room", 
    "Optometrist Room", 
    "Special Care Nursery",
    "Health Screening Unit",
    "Endoscopy Unit",
    "Daycare Ward",     "accident", "activities", "adl", "administration", "administrator",
    "affiliated", "amp", "and", "angiografi", "area", "assessment",
    "association", "asthma", "avenue", "barrier", "based", "basement",
    "bay", "bays", "beam", "bed", "bedded", "bereavement", "block",
    "bobath", "brachytherapy", "breastfeeding", "bulk", "bunker",
    "business", "c-arm", "cafeteria", "care", "cashier", "catheterization",
    "centre", "chair", "changing", "chemotherapy", "children", "clean",
    "clinic", "clinical", "clinics", "closet", "cold", "collection",
    "community", "cone", "consumable", "corridor", "couch", "counselling",
    "counter", "crrt", "cssd", "cssu", "cytotoxic", "day", "daycare",
    "dayward", "dental", "department", "desk", "dialysis", "dining",
    "dirty", "disposal", "district", "doctor's", "double", "double-bedded",
    "drug", "east", "echo", "ecg", "eeg", "endoscopy", "entrance",
    "equipment", "examination", "exam", "executive", "existing",
    "extended", "eye", "facilities", "facillities", "female", "fertility",
    "filing", "first", "floor", "focused", "four",
    "four-bedded", "function", "gamma", "general", "ground", "gym",
    "gymnasium", "haemodialysis", "hand", "hdu", "health", "high",
    "holding", "home", "hospice", "hospital", "hot", "icu", "imaging",
    "inhouse", "intensity", "intermediate", "international", "intra",
    "isolation", "it", "janitor", "janitor's", "joint", "knife",
    "laboratory", "laser", "ldr", "level", "lift", "linear", "lobby",
    "lounge", "lower", "lung", "major", "male", "management", "manager",
    "maternity", "medical", "medication", "medicine", "meeting",
    "mezzanine", "minor", "mobile", "multidisiplin", "negative",
    "nephrologist", "nerve", "new", "non-hepatitis", "nurse", "nursery",
    "nursing", "observation", "occupational", "off", "office", "oncology",
    "open", "operation", "optometrist", "optometry", "oral", "orthopaedic",
    "outpatient", "outsourced", "palliative", "park", "patient", "pediatric",
    "pharmacy", "physical", "physio", "physiotherapy", "physician's",
    "plastic", "play", "positive", "prep", "preparation", "pressure",
    "private", "procedure", "processing", "public", "quarantine",
    "radiology", "radiologist", "reception", "reconstitution", "record",
    "recovery", "refuge", "registration", "rehab", "rehabilitation",
    "relokasi", "reporting", "resusitasi", "retail", "robotic", "room",
    "rooms", "sample", "satellite", "scanner", "scn", "scope", "screening",
    "second", "serum", "server", "services", "settlement", "single",
    "single-bedded", "sleep", "sluice", "society", "soiled", "special",
    "specialist", "specialists", "speech", "staff", "station",
    "sterilization", "storage", "store", "stretcher", "study", "suite",
    "support", "surgery", "surgical", "table", "tea", "technical", "test",
    "theatre", "therapy", "third", "three", "three-bedded", "toilet",
    "tower", "traction", "training", "treadmill", "treatment", "triage",
    "trolley", "ultrasound", "unit", "utility", "vacant", "vascular",
    "vertigo", "vestibular", "vision", "visiting", "volunteer", "waiting",
    "ward", "wheelchair", "wing", "women's", "x-ray"
]

# Custom list of words/phrases to ALWAYS bold
BOLD_WORDS = [
    "Nama:",
    "No. K/P:",
    "Kelayakan:", "No. Lesen", "Tarikh Sah Lesen:", "(Borang 7):", "(Borang 4):", "diperakui", "diperakui dengan terma atau syarat"
]

def format_rich_text(text, italic_words, bold_words):
    """
    Converts plain text into a docxtpl RichText object.
    Applies Arial, Size 9, and targets specific words for Italics and/or Bold using Regex.
    """
    if not text:
        return RichText("")

    all_target_words = set(italic_words + bold_words)
    sorted_words = sorted(all_target_words, key=len, reverse=True)
    
    main_pattern = re.compile(r'(' + '|'.join(map(re.escape, sorted_words)) + r')', re.IGNORECASE)
    italic_pattern = re.compile(r'^(' + '|'.join(map(re.escape, italic_words)) + r')$', re.IGNORECASE)
    bold_pattern = re.compile(r'^(' + '|'.join(map(re.escape, bold_words)) + r')$', re.IGNORECASE)
    
    rt = RichText()
    FONT_KWARGS = {'font': 'Arial', 'size': 18} # 9 pt = 18 half-points
    
    parts = main_pattern.split(text)
    
    for part in parts:
        if not part:
            continue
            
        is_italic = bool(italic_pattern.match(part))
        is_bold = bool(bold_pattern.match(part))
        
        if is_italic and is_bold:
            rt.add(part, italic=True, bold=True, **FONT_KWARGS)
        elif is_italic:
            rt.add(part, italic=True, **FONT_KWARGS)
        elif is_bold:
            rt.add(part, bold=True, **FONT_KWARGS)
        else:
            rt.add(part, **FONT_KWARGS)
            
    return rt

def extract_data_from_docx(file):
    doc = docx.Document(file)
    
    raw_data = {
        "nama_premis": "", "no_fail": "", "alamat_premis": "",
        "no_siri_kemudahan": "", "no_siri_perkhidmatan": "", "t_s": "Tiada",
        "butiran_peluasan": "", "kapasiti": "", "no_lesen_tarikh": "",
        "jenis_entiti": "", "oydk": "", "tarikh_terima": "",
        "nama_entiti": "", "alamat_pendaftaran": "", "no_pendaftaran": "",
        "ulasan_ckaps": "", "cadangan_sekretariat": "", "keputusan": "",
        "pegawai_proses": ""
    }
    
    # ---------------------------------------------------------
    # 1. SWEEP TABLES (Extracts data from within table cells)
    # ---------------------------------------------------------
    capture_ckaps = False
    capture_sekretariat = False
    capture_keputusan = False
    capture_pegawai = False
    
    for table in doc.tables:
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and text != ":" and (not row_texts or text != row_texts[-1]):
                    row_texts.append(text)
            
            if not row_texts:
                continue
                
            if len(row_texts) > 1 and re.match(r'^\d+\.?$', row_texts[0].strip()):
                row_texts.pop(0)
                
            col_0_text = row_texts[0]
            col_1_text = row_texts[1] if len(row_texts) > 1 else ""
            combined_row_text = "\n".join(row_texts)

            # Capture Pegawai Proses if it was pushed to the next row
            if capture_pegawai:
                clean_text = combined_row_text.lstrip(":\n\r\t |")
                if clean_text:
                    raw_data["pegawai_proses"] = clean_text.split('\n')[0].strip()
                    capture_pegawai = False
                continue

            if "Disediakan oleh" in combined_row_text and not raw_data["pegawai_proses"]:
                capture_keputusan = False
                capture_sekretariat = False
                capture_ckaps = False
                
                text_after = re.split(r'(?i)Disediakan oleh\s*:?', combined_row_text)[-1]
                text_after = text_after.lstrip(":\n\r\t |")
                if text_after:
                    raw_data["pegawai_proses"] = text_after.split('\n')[0].strip()
                else:
                    capture_pegawai = True
                continue
            
            # --- 1. Nama Premis & No. Fail ---
            if "Nama Premis" in col_0_text:
                val = col_1_text if col_1_text else re.sub(r'(?i)Nama Premis\s*:?', '', combined_row_text).strip()
                match = re.search(r'(KKM\s*87/A2.*?)(?:\n|$)', val)
                if match:
                    raw_data["no_fail"] = match.group(1).strip()
                    raw_data["nama_premis"] = val.replace(match.group(1), "").strip().upper()
                else:
                    raw_data["nama_premis"] = val.strip().upper()
            
            # --- 2. Alamat Premis ---
            elif "Alamat Premis" in col_0_text:
                raw_data["alamat_premis"] = col_1_text if col_1_text else re.sub(r'(?i)Alamat Premis\s*:?', '', combined_row_text).strip()
                
            # --- 3. Jenis Kemudahan ---
            elif "Jenis Kemudahan" in col_0_text:
                match = re.search(r'No\.\s*Siri:\s*([A-Za-z0-9,\s&]+)', combined_row_text, re.IGNORECASE)
                if match:
                    raw_data["no_siri_kemudahan"] = match.group(1).replace(')', '').strip()
                if "Terma Atau Syarat:" in combined_row_text or "Terma atau Syarat:" in combined_row_text:
                    ts = re.split(r'(?i)Terma Atau Syarat:', combined_row_text)
                    if len(ts) > 1: raw_data["t_s"] = ts[1].strip()
                        
            # --- 4. Jenis Perkhidmatan ---
            elif "Jenis Perkhidmatan" in col_0_text:
                match = re.search(r'No\.\s*Siri:\s*([A-Za-z0-9,\s&]+)', combined_row_text, re.IGNORECASE)
                if match:
                    raw_data["no_siri_perkhidmatan"] = match.group(1).replace(')', '').strip()
                if "Terma Atau Syarat:" in combined_row_text or "Terma atau Syarat:" in combined_row_text:
                    ts = re.split(r'(?i)Terma Atau Syarat:', combined_row_text)
                    if len(ts) > 1: raw_data["t_s"] = ts[1].strip()
                        
            # --- 5. Butiran Peluasan ---
            elif "Butiran Peluasan" in col_0_text or "Butir-butir Peluasan" in col_0_text:
                raw_data["butiran_peluasan"] = col_1_text if col_1_text else re.sub(r'(?i)Butir(an|-butir) Peluasan( atau Pengubahan)?\s*:?', '', combined_row_text).strip()
                
            # --- 6. Kapasiti ---
            elif "Kapasiti" in col_0_text or "Saiz" in col_0_text:
                raw_data["kapasiti"] = col_1_text if col_1_text else re.sub(r'(?i)(Saiz/)?Kapasiti\s*:?', '', combined_row_text).strip()
                
            # --- 7. No Lesen & Tarikh ---
            elif "Tarikh Sah Lesen" in col_0_text or "Tarikh Luput Lesen" in col_0_text or "Butiran Lesen" in col_0_text or "No. Lesen" in col_0_text:
                val = col_1_text if col_1_text else re.sub(r'(?i)(Tarikh (Sah|Luput) Lesen|Butiran Lesen|No\. Lesen(\s*\(Borang\s*\d+\))?)\s*:?', '', combined_row_text).strip()
                if val: raw_data["no_lesen_tarikh"] += f"{val}\n"
                
            # --- 8. Jenis Usaha Niaga ---
            elif "Jenis Usaha Niaga" in col_0_text or "Jenis urus niaga" in col_0_text:
                raw_data["jenis_entiti"] = col_1_text if col_1_text else re.sub(r'(?i)Jenis (Usaha|urus) Nia`?ga\s*:?', '', combined_row_text).strip()
                
            # --- 9. Orang Yang Diberi Kuasa (OYDK) ---
            elif "Orang Yang Diberi Kuasa" in col_0_text or "OYDK" in col_0_text:
                raw_data["oydk"] = col_1_text if col_1_text else re.sub(r'(?i)(Orang Yang Diberi Kuasa(\s*\(OYDK\))?|OYDK)\s*:?', '', combined_row_text).strip()
                
            # --- 10. Tarikh Diterima ---
            elif "Tarikh Diterima" in col_0_text or "Tarikh Terima" in col_0_text:
                raw_data["tarikh_terima"] = col_1_text if col_1_text else re.sub(r'(?i)Tarikh D?iterima\s*:?', '', combined_row_text).strip()
                
            # --- 11. Nama Entiti ---
            elif "Nama Pertubuhan" in col_0_text:
                val = col_1_text if col_1_text else re.sub(r'(?i)Nama Pertubuhan( Perbadanan/ Perkongsian/Pertubuhan)?\s*:?', '', combined_row_text).strip()
                raw_data["nama_entiti"] = val.upper()
                
            # --- 12. Alamat Pendaftaran ---
            elif "Alamat Usaha Niaga" in col_0_text or "Nia`ga" in col_0_text:
                raw_data["alamat_pendaftaran"] = col_1_text if col_1_text else re.sub(r'(?i)Alamat Usaha Nia`?ga\s*:?', '', combined_row_text).strip()
                
            # --- 13. No Pendaftaran ---
            elif "No. Pendaftaran" in col_0_text:
                raw_data["no_pendaftaran"] = col_1_text if col_1_text else re.sub(r'(?i)No\. Pendaftaran( Syarikat)?\s*:?', '', combined_row_text).strip()
                
            # --- 14. Ulasan CKAPS ---
            elif "Ulasan Pelan Lantai" in col_0_text:
                capture_ckaps = True
                capture_sekretariat = False
                capture_keputusan = False
                continue
                
            # --- 15. Ulasan Sekretariat ---
            elif "Ulasan Sekretariat" in col_0_text or "Ulasan Sekreteriat" in col_0_text:
                capture_ckaps = False
                capture_sekretariat = True
                capture_keputusan = False
                
                val = col_1_text if col_1_text else re.sub(r'(?i)Ulasan Sekret[ea]riat:?', '', col_0_text).strip()
                if val: raw_data["cadangan_sekretariat"] += val + "\n"
                continue
                
            # --- 16. Keputusan ---
            elif "Keputusan Mesyuarat" in col_0_text or "KEPUTUSAN" in col_0_text or col_0_text.startswith("Keputusan:"):
                capture_ckaps = False
                capture_sekretariat = False
                capture_keputusan = True
                
                val = col_1_text if col_1_text else re.sub(r'(?i)Keputusan( Mesyuarat)?:?', '', col_0_text).strip()
                if val: raw_data["keputusan"] += val + "\n"
                continue

            # --- Multi-Row Active Appends ---
            if capture_ckaps:
                line = f"{col_0_text} {col_1_text}".strip()
                if line: raw_data["ulasan_ckaps"] += line + "\n"
            
            elif capture_sekretariat:
                if not re.match(r'^\d+\.$', col_0_text.strip()):
                    line = " ".join(row_texts).strip()
                    if line: raw_data["cadangan_sekretariat"] += line + "\n"
            
            elif capture_keputusan:
                line = " ".join(row_texts).strip()
                if line: raw_data["keputusan"] += line + "\n"

    # ---------------------------------------------------------
    # 2. SWEEP PARAGRAPHS (Extracts data typed OUTSIDE tables)
    # ---------------------------------------------------------
    capture_sekretariat_para = False
    capture_keputusan_para = False
    capture_pegawai_para = False
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        if capture_pegawai_para:
            clean_text = text.lstrip(":\n\r\t |")
            if clean_text:
                raw_data["pegawai_proses"] = clean_text.split('\n')[0].strip()
                capture_pegawai_para = False
            continue

        if "Disediakan oleh" in text and not raw_data["pegawai_proses"]:
            capture_sekretariat_para = False
            capture_keputusan_para = False
            
            text_after = re.split(r'(?i)Disediakan oleh\s*:?', text)[-1].lstrip(":\n\r\t |")
            if text_after:
                raw_data["pegawai_proses"] = text_after.split('\n')[0].strip()
            else:
                capture_pegawai_para = True
            continue
            
        if "Ulasan Sekretariat" in text or "Ulasan Sekreteriat" in text:
            capture_sekretariat_para = True
            capture_keputusan_para = False
            val = re.sub(r'(?i)Ulasan Sekret[ea]riat:?', '', text).strip()
            if val: raw_data["cadangan_sekretariat"] += val + "\n"
            continue
            
        if "Keputusan Mesyuarat" in text or "KEPUTUSAN" in text or text.startswith("Keputusan:"):
            capture_sekretariat_para = False
            capture_keputusan_para = True
            val = re.sub(r'(?i)Keputusan( Mesyuarat)?:?', '', text).strip()
            if val: raw_data["keputusan"] += val + "\n"
            continue
            
        if capture_sekretariat_para:
            if not re.match(r'^\d+\.$', text):
                raw_data["cadangan_sekretariat"] += text + "\n"
                
        if capture_keputusan_para:
            raw_data["keputusan"] += text + "\n"


    # ---------------------------------------------------------
    # 3. FORMATTING AND CLEANUP
    # ---------------------------------------------------------
    rich_data = {}
    for key, value in raw_data.items():
        rich_data[key] = format_rich_text(value.strip(), ENGLISH_WORDS, BOLD_WORDS) 
        
    rich_data["RAW_nama_premis"] = raw_data["nama_premis"].strip()

    return rich_data

def create_output_doc(data, template_path="Output.docx"):
    doc = DocxTemplate(template_path)
    doc.render(data)
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

# --- Streamlit UI ---
st.set_page_config(page_title="Buat Minit Borang 5 instant")
st.title("Instant Minit B5")
st.markdown("Muat naik fail permohonan, pastikan semua fail (.docx), Output perlu diupdate ikut tarikh JKP, hasil akan download dalam zip, perlu edit sikit2 format, OYDK, Ulasan CKAPS dan keputusan.")

uploaded_files = st.file_uploader("Muat naik fail Word", type=["docx"], accept_multiple_files=True)

if uploaded_files:
    if st.button("Jana Dokumen Output"):
        output_files = []
        
        for i, file in enumerate(uploaded_files, start=1):
            try:
                extracted_data = extract_data_from_docx(file)
                output_stream = create_output_doc(extracted_data)
                
                safe_name = f"Minit B5 {i:03d}.docx"
                
                output_files.append((safe_name, output_stream))
                st.success(f"Berjaya diproses: {safe_name} (dari {file.name})")
                
            except Exception as e:
                st.error(f"Ralat memproses {file.name}: {e}")

        if output_files:
            zip_buffer = io.BytesIO()
            
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for name, stream in output_files:
                    zip_file.writestr(name, stream.getvalue())
            
            zip_buffer.seek(0)
            
            st.download_button(
                label="📥 Muat Turun Semua (ZIP)",
                data=zip_buffer,
                file_name="Semua_Minit_B5.zip",
                mime="application/zip"
            )
